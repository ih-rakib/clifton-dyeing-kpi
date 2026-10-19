#!/usr/bin/env python3
"""
Clifton Apparels Ltd. - Dyeing & Finishing Performance Dashboard
================================================================
ETL builder.

Author : Ikramul Hasan Rakib
         Executive, Central IE Department, Clifton Group
(c) 2026 Ikramul Hasan Rakib. All rights reserved.

Reads the two live workbooks:
  * DYEING PERFORMANCE OVERVIEW.xlsx
  * FINISHING PERFORMANCE OVERVIEW.xlsx

and emits a single, compact data bundle consumed by the static dashboard.

Run manually:
    python scripts/build.py

The script is intentionally dependency-light (openpyxl only) so that it can
run on any machine in the plant and inside GitHub Actions.

Outputs
-------
dashboard/data/data.js     -> window.DASHBOARD_DATA = {...};   (works from file://)
dashboard/data/data.json   -> same payload, for API / tooling use
"""

from __future__ import annotations

import json
import math
import re
import sys
import traceback
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required.  Install with:  pip install -r requirements.txt")

ROOT = Path(__file__).resolve().parent.parent
DYEING_XLSX = ROOT / "DYEING PERFORMANCE OVERVIEW.xlsx"
FINISHING_XLSX = ROOT / "FINISHING PERFORMANCE OVERVIEW.xlsx"
OUT_DIR = ROOT / "dashboard" / "data"

# --------------------------------------------------------------------------
# Configuration - edit here to change targets / thresholds
# --------------------------------------------------------------------------
TARGETS = {
    "production_monthly_kg": 468_000,
    "oee": 85.0,
    "availability": 95.0,
    "performance": 90.0,
    "quality": 95.0,
    "rft": 90.0,
    "rft_bulk_to_bulk": 96.0,
    "rft_lab_to_bulk": 90.0,
    "topping_limit": 5.0,
    "redyeing_limit": 1.0,
    "hit_rate": 85.0,
    "finishing_oee": 70.0,
    "finishing_utilization": 90.0,
    "finishing_efficiency": 80.0,
    "finishing_quality": 98.0,
    "baseline_oee": 70.0,
    "baseline_availability": 75.0,
    "baseline_performance": 90.0,
    "baseline_quality": 95.0,
}

STANDARD_WORKING_DAYS = 26
SAMPLE_MACHINES = {"Mc-1", "Mc-2", "Mc-3", "Mc-25", "Mc-26"}

MONTH_ORDER = {
    m: i
    for i, m in enumerate(
        [
            "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
        ],
        1,
    )
}

# --------------------------------------------------------------------------
# Loss-time classification
# --------------------------------------------------------------------------
LOSS_RULES = [
    ("Heat-Set / Auto Stitch Delay", [
        "heat set", "heat-set", "heatset", "auto stitch", "autostitch", "auto-stitch",
        "heat setter",
    ]),
    ("Utilities (Steam / Gas / Power)", [
        "steam", "gas ", "gas.", "gas supply", "boiler", "pressure was low", "pdb",
        "power", "electric", "generator", "compressor", "water supply", "chiller",
    ]),
    ("Batch / Production Planning", [
        "no batch", "batch not ready", "not in plan", "in plan", "no plan", "plan",
        "order crisis", "no order", "program", "schedule change", "batch delay",
        "waiting for batch", "batch is not", "batch unavailable",
    ]),
    ("Maintenance / Planned Service", [
        "maintenance", "under maint", "servicing", "service", "preventive", "pm ",
        "overhaul", "scheduled",
    ]),
    ("Machine / Technical Failure", [
        "motor", "bearing", "inverter", "pump", "valve", "gasket", "reel", "nozzle",
        "leak", "damage", "breakdown", "sensor", "camera", "basket", "plc", "panel",
        "belt", "chain", "roller", "drum", "winch", "trolley", "temperature",
        "controller", "faulty", "problem", "repair", "jam", "block", "choke",
        "cloth guide", "seal", "pipe", "steam trap", "m/c stopped", "machine stopped",
        "m/c break", "cylinder", "pneumatic", "air line", "vfd", "encoder",
        "display", "module", "heater", "thermo",
    ]),
    ("Material / Handling", [
        "fabric unavail", "fabric not", "no fabric", "unload", "loading", "handling",
        "waiting for fabric", "t/f fabric", "t/f",
    ]),
    ("Quality / Process", [
        "shade", "rework", "reprocess", "re-dye", "redye", "topping", "rejection",
        "quality",
    ]),
]


def classify_loss(reason: str) -> str:
    if not reason:
        return "Other / Unclassified"
    r = str(reason).lower()
    for name, keys in LOSS_RULES:
        for k in keys:
            if k in r:
                return name
    return "Other / Unclassified"


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def num(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def opt_num(value):
    """Numeric value or None when the cell is empty."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def ratio(a: float, b: float):
    if not b:
        return None
    return a / b


def pct(a: float, b: float):
    r = ratio(a, b)
    return None if r is None else round(r * 100, 4)


def clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def canon_buyer(v) -> str:
    s = clean(v)
    if not s or s == "0":
        return "Unassigned"
    return s.upper()


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        s = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def month_key(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def month_label(d: date) -> str:
    return d.strftime("%b '%y")


def safe_round(v, nd=4):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    return round(v, nd)


# --------------------------------------------------------------------------
# Header map helper
# --------------------------------------------------------------------------
def norm(s: str) -> str:
    s = clean(s).lower()
    s = re.sub(r"[\r\n]", " ", s)
    s = re.sub(r"[^a-z0-9%/()]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


class Row(dict):
    """Row accessor tolerant to header whitespace/case variations."""

    def __init__(self, mapping):
        super().__init__(mapping)

    def get(self, *names, default=None):
        for n in names:
            k = norm(n)
            if k in self:
                return self[k]
        return default


def build_reader(path: Path, sheet: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Sheet {sheet!r} not found in {path.name}")
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    keys = [norm(h) if h is not None else f"__col{i}" for i, h in enumerate(header)]
    return wb, it, keys


def make_row(values, keys):
    d = {}
    for k, v in zip(keys, values):
        # keep the first occurrence of duplicated header names
        if k not in d:
            d[k] = v
    return Row(d)


# ==========================================================================
# DYEING
# ==========================================================================
def load_dyeing():
    wb, it, keys = build_reader(DYEING_XLSX, "Production Report")

    monthly = defaultdict(lambda: defaultdict(float))
    daily = defaultdict(lambda: defaultdict(float))
    machine = defaultdict(lambda: defaultdict(float))
    machine_month = defaultdict(lambda: defaultdict(float))
    machine_day = defaultdict(lambda: defaultdict(float))
    shift_month = defaultdict(lambda: defaultdict(float))
    mix = defaultdict(lambda: defaultdict(Counter))
    loss_reason = defaultdict(lambda: defaultdict(lambda: {"hours": 0.0, "count": 0}))
    dates_by_month = defaultdict(set)
    machines_by_month = defaultdict(set)
    records_by_month = Counter()
    blank_rows = 0
    row_perf = defaultdict(lambda: [0.0, 0])
    row_oee = defaultdict(lambda: [0.0, 0])
    row_avail = defaultdict(lambda: [0.0, 0])
    row_qual = defaultdict(lambda: [0.0, 0])

    for values in it:
        if values is None:
            continue
        r = make_row(values, keys)
        d = to_date(r.get("date"))
        if d is None:
            blank_rows += 1
            continue

        mk = month_key(d)
        mname = clean(r.get("machine no.")) or "Unknown"
        shift = clean(r.get("shift")) or "-"
        mtype = clean(r.get("machine type")) or ("Sample" if mname in SAMPLE_MACHINES else "Bulk")

        prod = num(r.get("production/shift (kg)", "production/shift (kg) "))
        target = num(r.get("target/ shift (kg)", "target/shift (kg)", "target / shift (kg)"))
        load = num(r.get("loading capacity"))
        wh = num(r.get("working hour"))
        loss_h = num(r.get("loss time (hour)", "loss time (hour) "))
        op_h = opt_num(r.get("operating time (hour)", "operating time (hour) "))
        if op_h is None:
            op_h = max(wh - loss_h, 0.0)
        qpp = num(r.get("quality pass production (kg)"))
        if not qpp:
            qpp = max(target - num(r.get("topping")), 0.0)
        topping = num(r.get("topping"))
        redye = num(r.get("redyeing"))
        loss_kg = num(r.get("loss production(kg) /day"))

        dates_by_month[mk].add(d)
        if mtype != "Sample":
            machines_by_month[mk].add(mname)
        records_by_month[mk] += 1

        for bucket in (monthly[mk], daily[d], machine[mname], machine_month[(mname, mk)],
                       machine_day[(mname, d)], shift_month[(shift, mk)]):
            bucket["production"] += prod
            bucket["target"] += target
            bucket["loading_capacity"] += load
            bucket["working_hours"] += wh
            bucket["operating_hours"] += op_h
            bucket["loss_hours"] += loss_h
            bucket["quality_pass"] += qpp
            bucket["topping"] += topping
            bucket["redyeing"] += redye
            bucket["loss_kg"] += loss_kg
            bucket["records"] += 1

        monthly[mk]["sample_production"] += prod if mtype == "Sample" else 0.0
        monthly[mk]["bulk_production"] += prod if mtype != "Sample" else 0.0
        monthly[mk]["sample_records"] += 1 if mtype == "Sample" else 0

        # machine type flag
        if mtype == "Sample":
            machine[mname]["is_sample"] = 1.0

        # production mix
        mix[mk]["buyer"][canon_buyer(r.get("buyer"))] += prod
        mix[mk]["fabric"][clean(r.get("fabric type")) or "Unspecified"] += prod
        mix[mk]["color"][clean(r.get("color type")) or "Unspecified"] += prod
        mix[mk]["gsm"][clean(r.get("type of gsm")) or "Unspecified"] += prod
        mix[mk]["dia"][clean(r.get("type of dia")) or "Unspecified"] += prod

        # loss time
        reason = clean(r.get("remarks due to lose time", "remarks due to loss time"))
        if reason or loss_h:
            key = reason if reason else "Unspecified delay"
            loss_reason[mk][key]["hours"] += loss_h
            loss_reason[mk][key]["count"] += 1

        # row level ratios (report-style averages)
        p_row = opt_num(r.get("performance % (loading capacity vs production)"))
        o_row = opt_num(r.get("oee %"))
        a_row = opt_num(r.get("availability %"))
        q_row = opt_num(r.get("quality %"))
        if p_row is not None:
            row_perf[mk][0] += p_row
            row_perf[mk][1] += 1
        if o_row is not None:
            row_oee[mk][0] += o_row
            row_oee[mk][1] += 1
        if a_row is not None:
            row_avail[mk][0] += a_row
            row_avail[mk][1] += 1
        if q_row is not None:
            row_qual[mk][0] += q_row
            row_qual[mk][1] += 1

    wb.close()

    # ---- finalise monthly -------------------------------------------------
    months = []
    for mk in sorted(monthly.keys()):
        a = monthly[mk]
        year, mon = map(int, mk.split("-"))
        d0 = date(year, mon, 1)
        working_days = len(dates_by_month[mk])
        partial = False
        if working_days:
            last = max(dates_by_month[mk])
            next_month = date(year + (mon == 12), (mon % 12) + 1, 1)
            last_overall = max(dates_by_month[max(dates_by_month)])
            if d0.year == last_overall.year and d0.month == last_overall.month:
                partial = (next_month - last).days > 3

        availability = pct(a["operating_hours"], a["working_hours"])
        performance_w = pct(a["production"], a["loading_capacity"])
        quality = pct(a["quality_pass"], a["target"])
        oee_w = (
            round(availability * performance_w * quality / 10000, 4)
            if None not in (availability, performance_w, quality)
            else None
        )
        # report-style: simple average of row level ratios
        # (computed after the loop below from machine_day/month sums is not equal;
        #  we pre-compute using the exact row averages)
        goal = TARGETS["production_monthly_kg"]
        prorated_goal = goal * working_days / STANDARD_WORKING_DAYS if working_days else goal
        months.append({
            "key": mk,
            "label": month_label(d0),
            "monthName": d0.strftime("%B"),
            "year": year,
            "month": mon,
            "partial": partial,
            "workingDays": working_days,
            "activeMachines": len(machines_by_month[mk]),
            "records": int(a["records"]),
            "sampleRecords": int(a["sample_records"]),
            "workingHours": safe_round(a["working_hours"], 2),
            "operatingHours": safe_round(a["operating_hours"], 2),
            "lossHours": safe_round(a["loss_hours"], 2),
            "production": safe_round(a["production"], 2),
            "bulkProduction": safe_round(a["bulk_production"], 2),
            "sampleProduction": safe_round(a["sample_production"], 2),
            "plannedOutput": safe_round(a["target"], 2),
            "loadingCapacity": safe_round(a["loading_capacity"], 2),
            "qualityPass": safe_round(a["quality_pass"], 2),
            "topping": safe_round(a["topping"], 2),
            "redyeing": safe_round(a["redyeing"], 2),
            "lossKg": safe_round(a["loss_kg"], 2),
            "availability": safe_round(availability, 2),
            "performance": safe_round(performance_w, 2),
            "quality": safe_round(quality, 2),
            "oee": safe_round(oee_w, 2),
            "hitRate": safe_round(pct(a["production"], goal), 2),
            "capacityUtil": safe_round(pct(a["production"], a["loading_capacity"]), 2),
            "toppingPct": safe_round(pct(a["topping"], a["target"]), 2),
            "redyeingPct": safe_round(pct(a["redyeing"], a["target"]), 2),
            "reworkPct": safe_round(pct(a["topping"] + a["redyeing"], a["target"]), 2),
            "goal": goal,
            "proratedGoal": safe_round(prorated_goal, 0),
            "hitRateProrated": safe_round(pct(a["production"], prorated_goal), 2),
            "avgDailyOutput": safe_round(a["production"] / working_days, 1) if working_days else None,
            "avgDailyLossKg": safe_round(a["loss_kg"] / working_days, 1) if working_days else None,
        })

    # ---- report-style monthly ratios (row average) -----------------------
    report_style = {}
    for mk in set(list(row_perf) + list(row_oee) + list(row_avail) + list(row_qual)):
        report_style[mk] = {
            "performance": safe_round(row_perf[mk][0] / row_perf[mk][1] * 100, 2) if row_perf[mk][1] else None,
            "oee": safe_round(row_oee[mk][0] / row_oee[mk][1] * 100, 2) if row_oee[mk][1] else None,
            "availability": safe_round(row_avail[mk][0] / row_avail[mk][1] * 100, 2) if row_avail[mk][1] else None,
            "quality": safe_round(row_qual[mk][0] / row_qual[mk][1] * 100, 2) if row_qual[mk][1] else None,
        }

    for m in months:
        rs = report_style.get(m["key"], {})
        m["performanceReport"] = rs.get("performance")
        m["oeeReport"] = rs.get("oee")
        m["availabilityReport"] = rs.get("availability")
        m["qualityReport"] = rs.get("quality")

    # ---- daily ------------------------------------------------------------
    daily_rows = []
    for d in sorted(daily.keys()):
        a = daily[d]
        mk = month_key(d)
        daily_rows.append({
            "date": d.isoformat(),
            "month": mk,
            "production": safe_round(a["production"], 1),
            "plannedOutput": safe_round(a["target"], 1),
            "loadingCapacity": safe_round(a["loading_capacity"], 1),
            "workingHours": safe_round(a["working_hours"], 1),
            "operatingHours": safe_round(a["operating_hours"], 1),
            "lossHours": safe_round(a["loss_hours"], 1),
            "lossKg": safe_round(a["loss_kg"], 1),
            "availability": safe_round(pct(a["operating_hours"], a["working_hours"]), 2),
            "performance": safe_round(pct(a["production"], a["loading_capacity"]), 2),
            "quality": safe_round(pct(a["quality_pass"], a["target"]), 2),
            "hitRate": safe_round(pct(a["production"], a["target"]), 2),
            "records": int(a["records"]),
        })
        gr = daily_rows[-1]
        if None not in (gr["availability"], gr["performance"], gr["quality"]):
            gr["oee"] = round(gr["availability"] * gr["performance"] * gr["quality"] / 10000, 2)
        else:
            gr["oee"] = None

    # ---- machines ---------------------------------------------------------
    machine_rows = []
    for name, a in machine.items():
        days_active = len({d for (mn, d), v in machine_day.items() if mn == name and v["production"] > 0})
        availability = pct(a["operating_hours"], a["working_hours"])
        performance = pct(a["production"], a["loading_capacity"])
        quality = pct(a["quality_pass"], a["target"])
        oee = (
            round(availability * performance * quality / 10000, 2)
            if None not in (availability, performance, quality)
            else None
        )
        is_sample = name in SAMPLE_MACHINES
        machine_rows.append({
            "machine": name,
            "isSample": is_sample,
            "production": safe_round(a["production"], 1),
            "loadingCapacity": safe_round(a["loading_capacity"], 1),
            "plannedOutput": safe_round(a["target"], 1),
            "workingHours": safe_round(a["working_hours"], 1),
            "operatingHours": safe_round(a["operating_hours"], 1),
            "lossHours": safe_round(a["loss_hours"], 1),
            "qualityPass": safe_round(a["quality_pass"], 1),
            "topping": safe_round(a["topping"], 1),
            "redyeing": safe_round(a["redyeing"], 1),
            "availability": safe_round(availability, 2),
            "performance": safe_round(performance, 2),
            "quality": safe_round(quality, 2),
            "oee": safe_round(oee, 2),
            "hitRate": safe_round(pct(a["production"], a["target"]), 2),
            "records": int(a["records"]),
            "daysActive": days_active,
            "avgDailyOutput": safe_round(a["production"] / days_active, 1) if days_active else None,
        })
    machine_rows.sort(key=lambda x: (x["isSample"], -(x["oee"] if x["oee"] is not None else -1)))

    # ---- machine monthly series ------------------------------------------
    machine_monthly = defaultdict(list)
    for (name, mk), a in sorted(machine_month.items()):
        availability = pct(a["operating_hours"], a["working_hours"])
        performance = pct(a["production"], a["loading_capacity"])
        quality = pct(a["quality_pass"], a["target"])
        oee = (
            round(availability * performance * quality / 10000, 2)
            if None not in (availability, performance, quality)
            else None
        )
        days = len({d for (mn, d), v in machine_day.items() if mn == name and month_key(d) == mk and v["production"] > 0})
        machine_monthly[name].append({
            "month": mk,
            "production": safe_round(a["production"], 1),
            "loadingCapacity": safe_round(a["loading_capacity"], 1),
            "plannedOutput": safe_round(a["target"], 1),
            "workingHours": safe_round(a["working_hours"], 1),
            "operatingHours": safe_round(a["operating_hours"], 1),
            "lossHours": safe_round(a["loss_hours"], 1),
            "qualityPass": safe_round(a["quality_pass"], 1),
            "availability": safe_round(availability, 2),
            "performance": safe_round(performance, 2),
            "quality": safe_round(quality, 2),
            "oee": safe_round(oee, 2),
            "records": int(a["records"]),
            "days": days,
        })

    # ---- heatmap (machine x day) -----------------------------------------
    heatmap = {}
    for mk in [m["key"] for m in months]:
        days = sorted(dates_by_month[mk])
        names = sorted({mn for (mn, d) in machine_day if month_key(d) == mk})
        oee_matrix = []
        prod_matrix = []
        for name in names:
            oee_row, prod_row = [], []
            for d in days:
                a = machine_day.get((name, d))
                if not a:
                    oee_row.append(None)
                    prod_row.append(None)
                    continue
                availability = pct(a["operating_hours"], a["working_hours"])
                performance = pct(a["production"], a["loading_capacity"])
                quality = pct(a["quality_pass"], a["target"])
                o = (
                    round(availability * performance * quality / 10000, 1)
                    if None not in (availability, performance, quality)
                    else None
                )
                oee_row.append(o)
                prod_row.append(round(a["production"], 1))
            oee_matrix.append(oee_row)
            prod_matrix.append(prod_row)
        heatmap[mk] = {
            "dates": [d.isoformat() for d in days],
            "machines": names,
            "oee": oee_matrix,
            "production": prod_matrix,
        }

    # ---- shift ------------------------------------------------------------
    shifts = {}
    for (shift, mk), a in shift_month.items():
        availability = pct(a["operating_hours"], a["working_hours"])
        performance = pct(a["production"], a["loading_capacity"])
        quality = pct(a["quality_pass"], a["target"])
        oee = (
            round(availability * performance * quality / 10000, 2)
            if None not in (availability, performance, quality)
            else None
        )
        shifts.setdefault(mk, {})[shift] = {
            "production": safe_round(a["production"], 1),
            "availability": safe_round(availability, 2),
            "performance": safe_round(performance, 2),
            "quality": safe_round(quality, 2),
            "oee": safe_round(oee, 2),
            "lossHours": safe_round(a["loss_hours"], 1),
            "records": int(a["records"]),
        }

    # ---- loss -------------------------------------------------------------
    loss_monthly = {}
    for mk, reasons in loss_reason.items():
        cats = defaultdict(lambda: {"hours": 0.0, "count": 0})
        top = []
        for reason, v in reasons.items():
            cat = classify_loss(reason)
            cats[cat]["hours"] += v["hours"]
            cats[cat]["count"] += v["count"]
            if v["hours"] > 0:
                top.append({
                    "reason": reason,
                    "hours": round(v["hours"], 2),
                    "count": v["count"],
                    "category": cat,
                })
        total = sum(c["hours"] for c in cats.values())
        cat_list = [
            {
                "category": k,
                "hours": round(v["hours"], 2),
                "count": v["count"],
                "share": safe_round(v["hours"] / total * 100, 2) if total else 0,
            }
            for k, v in cats.items()
            if v["hours"] > 0
        ]
        cat_list.sort(key=lambda x: -x["hours"])
        top.sort(key=lambda x: -x["hours"])
        loss_monthly[mk] = {
            "totalHours": round(total, 2),
            "categories": cat_list,
            "topReasons": top[:25],
        }

    # loss by machine (all time + per month handled in UI by machineMonthly)
    loss_by_machine = sorted(
        [
            {"machine": m["machine"], "lossHours": m["lossHours"], "oee": m["oee"],
             "production": m["production"]}
            for m in machine_rows
            if not m["isSample"]
        ],
        key=lambda x: -(x["lossHours"] or 0),
    )

    # ---- mix --------------------------------------------------------------
    mix_out = {}
    for mk, groups in mix.items():
        block = {}
        for g, counter in groups.items():
            items = counter.most_common()
            if g in ("fabric", "color"):
                top = items[:12]
                other = sum(v for _, v in items[12:])
                arr = [[k or "Unspecified", round(v, 1)] for k, v in top]
                if other > 0:
                    arr.append(["Other", round(other, 1)])
            else:
                arr = [[k or "Unspecified", round(v, 1)] for k, v in items[:20]]
            total = sum(v for _, v in items)
            block[g] = {"items": arr, "total": round(total, 1)}
        mix_out[mk] = block

    # ---- RFT --------------------------------------------------------------
    rft = load_rft()

    return {
        "months": months,
        "daily": daily_rows,
        "machines": machine_rows,
        "machineMonthly": dict(machine_monthly),
        "heatmap": heatmap,
        "shifts": shifts,
        "loss": {
            "monthly": loss_monthly,
            "byMachine": loss_by_machine,
        },
        "mix": mix_out,
        "rft": rft,
        "blankRows": blank_rows,
    }


def load_rft():
    wb, it, keys = build_reader(DYEING_XLSX, "RFT Report")
    monthly = defaultdict(lambda: defaultdict(float))
    daily = defaultdict(lambda: defaultdict(float))
    has_data = False
    for values in it:
        r = make_row(values, keys)
        d = to_date(r.get("date"))
        kind = clean(r.get("type of rft"))
        if d is None or kind not in ("Bulk to Bulk", "Lab to Bulk"):
            continue
        batches = opt_num(r.get("number of batch"))
        passed = opt_num(r.get("number of batch rft pass"))
        failed = opt_num(r.get("number of batch rft fail"))
        if batches is None:
            continue
        has_data = True
        mk = month_key(d)
        slot = "b2b" if kind == "Bulk to Bulk" else "l2b"
        for bucket in (monthly[mk], daily[d]):
            bucket["total"] += batches
            bucket[slot] += batches
            bucket["pass"] += passed or 0
            bucket[slot + "_pass"] += passed or 0
            bucket["fail"] += failed or 0
            bucket[slot + "_fail"] += failed or 0
    wb.close()

    months = []
    for mk in sorted(monthly.keys()):
        a = monthly[mk]
        months.append({
            "month": mk,
            "total": int(a["total"]),
            "pass": int(a["pass"]),
            "fail": int(a["fail"]),
            "rate": safe_round(pct(a["pass"], a["total"]), 2),
            "b2b": {
                "total": int(a["b2b"]),
                "pass": int(a["b2b_pass"]),
                "rate": safe_round(pct(a["b2b_pass"], a["b2b"]), 2),
            },
            "l2b": {
                "total": int(a["l2b"]),
                "pass": int(a["l2b_pass"]),
                "rate": safe_round(pct(a["l2b_pass"], a["l2b"]), 2),
            },
        })

    daily_rows = []
    for d in sorted(daily.keys()):
        a = daily[d]
        daily_rows.append({
            "date": d.isoformat(),
            "month": month_key(d),
            "total": int(a["total"]),
            "pass": int(a["pass"]),
            "rate": safe_round(pct(a["pass"], a["total"]), 2),
            "b2b": int(a["b2b"]),
            "b2bPass": int(a["b2b_pass"]),
            "l2b": int(a["l2b"]),
            "l2bPass": int(a["l2b_pass"]),
        })

    return {"monthly": months, "daily": daily_rows, "available": has_data}


# ==========================================================================
# FINISHING
# ==========================================================================
def load_finishing():
    wb, it, keys = build_reader(FINISHING_XLSX, "Database")

    monthly = defaultdict(lambda: defaultdict(float))
    daily = defaultdict(lambda: defaultdict(float))
    machine = defaultdict(lambda: defaultdict(float))
    machine_month = defaultdict(lambda: defaultdict(float))
    fabric_daily = defaultdict(lambda: defaultdict(float))
    loss_reason = defaultdict(lambda: defaultdict(lambda: {"hours": 0.0, "count": 0}))
    records = Counter()
    dates_by_month = defaultdict(set)
    machines_by_month = defaultdict(set)
    machine_days = defaultdict(set)

    for values in it:
        r = make_row(values, keys)
        d = to_date(r.get("date"))
        if d is None:
            continue
        mk = month_key(d)
        mname = clean(r.get("machine name")) or "Unknown"
        fabric = clean(r.get("fabric type")) or "Unspecified"
        cap = num(r.get("capacity (kg)"))
        act = num(r.get("actual production (kg)"))
        rep = num(r.get("re-process (kg)"))
        qpp = num(r.get("quality pass production"))
        tt = opt_num(r.get("target time (hours)"))
        dt = num(r.get("downtime (hours)"))
        rt = opt_num(r.get("running time (hr)"))
        if rt is None and tt is not None:
            rt = max(tt - dt, 0.0)
        if not qpp and act:
            qpp = max(act - rep, 0.0)

        records[mk] += 1
        dates_by_month[mk].add(d)
        if act or cap or (tt or 0) > 0:
            machines_by_month[mk].add(mname)
        machine_days[mname].add(d)

        for bucket in (monthly[mk], daily[d], machine[mname], machine_month[(mname, mk)]):
            bucket["production"] += act
            bucket["capacity"] += cap
            bucket["reprocess"] += rep
            bucket["quality_pass"] += qpp
            bucket["target_time"] += tt or 0.0
            bucket["downtime"] += dt
            bucket["running_time"] += rt or 0.0
            bucket["records"] += 1
            if tt is not None:
                bucket["target_time_valid"] += tt
                bucket["running_time_valid"] += rt or 0.0
            if cap:
                bucket["capacity_valid"] += cap
                bucket["production_valid"] += act

        fabric_daily[(d, fabric)]["production"] += act
        fabric_daily[(d, fabric)]["capacity"] += cap
        fabric_daily[(d, fabric)]["downtime"] += dt
        fabric_daily[(d, fabric)]["target_time"] += tt or 0.0
        fabric_daily[(d, fabric)]["running_time"] += rt or 0.0
        fabric_daily[(d, fabric)]["quality_pass"] += qpp

        reason = clean(r.get("delay reason"))
        if reason and dt:
            loss_reason[mk][reason]["hours"] += dt
            loss_reason[mk][reason]["count"] += 1

    wb.close()

    def summarise(a):
        util = pct(a.get("running_time_valid", 0), a.get("target_time_valid", 0))
        eff = pct(a.get("production_valid", 0), a.get("capacity_valid", 0))
        qual = pct(a.get("quality_pass", 0), a.get("production", 0))
        oee = (
            round(util * eff * qual / 10000, 2)
            if None not in (util, eff, qual)
            else None
        )
        return {
            "production": safe_round(a.get("production", 0), 1),
            "capacity": safe_round(a.get("capacity", 0), 1),
            "reprocess": safe_round(a.get("reprocess", 0), 1),
            "qualityPass": safe_round(a.get("quality_pass", 0), 1),
            "targetTime": safe_round(a.get("target_time", 0), 2),
            "downtime": safe_round(a.get("downtime", 0), 2),
            "runningTime": safe_round(a.get("running_time", 0), 2),
            "utilization": safe_round(util, 2),
            "efficiency": safe_round(eff, 2),
            "quality": safe_round(qual, 2),
            "oee": safe_round(oee, 2),
            "reworkPct": safe_round(pct(a.get("reprocess", 0), a.get("production", 0)), 2),
            "records": int(a.get("records", 0)),
        }

    months = []
    for mk in sorted(monthly.keys()):
        year, mon = map(int, mk.split("-"))
        d0 = date(year, mon, 1)
        block = summarise(monthly[mk])
        block.update({
            "key": mk,
            "label": month_label(d0),
            "monthName": d0.strftime("%B"),
            "year": year,
            "month": mon,
            "workingDays": len(dates_by_month[mk]),
            "machines": len(machines_by_month[mk]),
        })
        months.append(block)

    daily_rows = []
    for d in sorted(daily.keys()):
        a = daily[d]
        block = summarise(a)
        block.update({"date": d.isoformat(), "month": month_key(d)})
        daily_rows.append(block)

    # per-day fabric split
    fab_map = defaultdict(dict)
    for (d, fabric), a in fabric_daily.items():
        fab_map[month_key(d)].setdefault(fabric, defaultdict(float))
        for k, v in a.items():
            fab_map[month_key(d)][fabric][k] += v
    fabric_monthly = {}
    for mk, fabrics in fab_map.items():
        fabric_monthly[mk] = {
            name: summarise(dict(vals)) for name, vals in fabrics.items()
        }

    machines = []
    for name, a in machine.items():
        block = summarise(a)
        block["machine"] = name
        block["daysActive"] = len(machine_days[name])
        machines.append(block)
    machines.sort(key=lambda x: -(x["production"] or 0))

    machine_monthly = {}
    for (name, mk), a in machine_month.items():
        machine_monthly.setdefault(name, {})[mk] = summarise(a)

    loss_monthly = {}
    for mk, reasons in loss_reason.items():
        cats = defaultdict(lambda: {"hours": 0.0, "count": 0})
        top = []
        for reason, v in reasons.items():
            cat = classify_loss(reason)
            cats[cat]["hours"] += v["hours"]
            cats[cat]["count"] += v["count"]
            top.append({"reason": reason, "hours": round(v["hours"], 2),
                        "count": v["count"], "category": cat})
        total = sum(c["hours"] for c in cats.values())
        cat_list = sorted(
            [{"category": k, "hours": round(v["hours"], 2), "count": v["count"],
              "share": safe_round(v["hours"] / total * 100, 2) if total else 0}
             for k, v in cats.items() if v["hours"] > 0],
            key=lambda x: -x["hours"],
        )
        top.sort(key=lambda x: -x["hours"])
        loss_monthly[mk] = {"totalHours": round(total, 2), "categories": cat_list,
                            "topReasons": top[:25]}

    # downtime by machine per month
    downtime_machine = {}
    for (name, mk), a in machine_month.items():
        downtime_machine.setdefault(mk, []).append({
            "machine": name,
            "downtime": safe_round(a["downtime"], 2),
            "production": safe_round(a["production"], 1),
        })
    for mk in downtime_machine:
        downtime_machine[mk].sort(key=lambda x: -x["downtime"])

    summary_daily = load_finishing_summary(max(daily.keys()) if daily else None)
    finished_by_month = defaultdict(float)
    for row in summary_daily:
        finished_by_month[row["month"]] += row["finished"]

    rework_total = sum(a["reprocess"] for a in monthly.values())

    for m in months:
        m["finished"] = safe_round(finished_by_month.get(m["key"], 0), 1) or None

    return {
        "months": months,
        "daily": daily_rows,
        "machines": machines,
        "machineMonthly": machine_monthly,
        "fabric": fabric_monthly,
        "loss": loss_monthly,
        "downtimeByMachine": downtime_machine,
        "summaryDaily": summary_daily,
        "finishedByMonth": {k: safe_round(v, 1) for k, v in finished_by_month.items()},
        "reworkTracked": rework_total > 0,
    }


def load_finishing_summary(max_date=None):
    """Daily finished-goods output and dyeing unload from the Summary sheet."""
    wb = openpyxl.load_workbook(FINISHING_XLSX, read_only=True, data_only=True)
    if "Summary" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Summary"]
    out = []
    for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        d = to_date(row[1])
        if d is None:
            continue
        if max_date is not None and d > max_date:
            continue
        finished = opt_num(row[2])
        dyeing = opt_num(row[3])
        if finished is None and dyeing is None:
            continue
        out.append({
            "date": d.isoformat(),
            "month": month_key(d),
            "finished": safe_round(finished or 0, 1),
            "dyeingUnload": safe_round(dyeing or 0, 1),
        })
    wb.close()
    out.sort(key=lambda x: x["date"])
    return out


# ==========================================================================
# INSIGHTS
# ==========================================================================
def build_insights(dyeing, finishing):
    months = dyeing["months"]
    rft_monthly = {r["month"]: r for r in dyeing["rft"]["monthly"]}
    fin_monthly = {m["key"]: m for m in finishing["months"]}
    insights = {}

    for idx, m in enumerate(months):
        mk = m["key"]
        prev = months[idx - 1] if idx > 0 else None
        items = []

        # Production vs goal
        if m["hitRate"] is not None:
            gap = m["production"] - m["goal"]
            if m["partial"]:
                rate = m["hitRateProrated"] or 0
                items.append({
                    "type": "info",
                    "title": "Month-to-date output",
                    "text": f"{m['production']:,.0f} kg produced across {m['workingDays']} working days — "
                            f"{rate:.1f}% of the prorated {m['goal']:,.0f} kg goal ({(m['proratedGoal'] or 0):,.0f} kg).",
                })
            elif m["hitRate"] >= 95:
                items.append({
                    "type": "positive",
                    "title": "Production goal achieved",
                    "text": f"Output of {m['production']:,.0f} kg reached {m['hitRate']:.1f}% of the "
                            f"{m['goal']:,.0f} kg monthly goal ({gap:+,.0f} kg).",
                })
            elif m["hitRate"] >= 85:
                items.append({
                    "type": "warning",
                    "title": "Production near goal",
                    "text": f"Output of {m['production']:,.0f} kg reached {m['hitRate']:.1f}% of goal — "
                            f"{abs(gap):,.0f} kg short of the {m['goal']:,.0f} kg target.",
                })
            else:
                items.append({
                    "type": "critical",
                    "title": "Production below goal",
                    "text": f"Output of {m['production']:,.0f} kg reached only {m['hitRate']:.1f}% of the "
                            f"{m['goal']:,.0f} kg target ({abs(gap):,.0f} kg shortfall).",
                })

        # OEE
        if m["oee"] is not None:
            delta = (m["oee"] - prev["oee"]) if prev and prev["oee"] is not None else None
            trend = ""
            if delta is not None:
                trend = f" ({delta:+.1f} pp vs {prev['label']})"
            if m["oee"] >= TARGETS["oee"]:
                items.append({"type": "positive", "title": "OEE above target",
                              "text": f"OEE closed at {m['oee']:.1f}% against the {TARGETS['oee']:.0f}% target{trend}."})
            elif m["oee"] >= TARGETS["oee"] - 5:
                items.append({"type": "warning", "title": "OEE slightly below target",
                              "text": f"OEE was {m['oee']:.1f}% vs a {TARGETS['oee']:.0f}% target{trend}."})
            else:
                items.append({"type": "critical", "title": "OEE below target",
                              "text": f"OEE fell to {m['oee']:.1f}%, {TARGETS['oee'] - m['oee']:.1f} pp below target{trend}."})

        # OEE loss driver
        if None not in (m["availability"], m["performance"], m["quality"]):
            components = [
                ("Availability", m["availability"], TARGETS["availability"]),
                ("Performance", m["performance"], TARGETS["performance"]),
                ("Quality", m["quality"], TARGETS["quality"]),
            ]
            worst = min(components, key=lambda c: c[1] - c[2])
            gap = worst[2] - worst[1]
            if gap > 0.5:
                items.append({
                    "type": "info",
                    "title": f"Biggest OEE gap: {worst[0]}",
                    "text": f"{worst[0]} at {worst[1]:.1f}% is {gap:.1f} pp under its {worst[2]:.0f}% target — "
                            f"the single largest recoverable loss in the OEE stack.",
                })

        # Availability
        if m["availability"] is not None and m["availability"] >= TARGETS["availability"]:
            items.append({"type": "positive", "title": "Availability strong",
                          "text": f"Availability held at {m['availability']:.1f}% with {m['lossHours']:,.1f} loss hours "
                                  f"across {m['workingDays']} working days."})

        # RFT
        r = rft_monthly.get(mk)
        if r and r["rate"] is not None:
            if r["rate"] >= TARGETS["rft"]:
                items.append({"type": "positive", "title": "Right-first-time above target",
                              "text": f"RFT reached {r['rate']:.1f}% ({r['pass']:,}/{r['total']:,} batches) against a "
                                      f"{TARGETS['rft']:.0f}% target."})
            else:
                items.append({"type": "warning", "title": "Right-first-time below target",
                              "text": f"RFT was {r['rate']:.1f}% ({r['pass']:,}/{r['total']:,} batches), "
                                      f"{TARGETS['rft'] - r['rate']:.1f} pp below the {TARGETS['rft']:.0f}% target."})

        # Rework
        if m["toppingPct"] is not None:
            if m["toppingPct"] > TARGETS["topping_limit"]:
                items.append({"type": "warning", "title": "Topping above limit",
                              "text": f"Topping stood at {m['toppingPct']:.2f}% ({m['topping']:,.0f} kg), above the "
                                      f"{TARGETS['topping_limit']:.0f}% ceiling."})
            else:
                items.append({"type": "positive", "title": "Topping within limit",
                              "text": f"Topping was {m['toppingPct']:.2f}% ({m['topping']:,.0f} kg) against a "
                                      f"{TARGETS['topping_limit']:.0f}% ceiling."})
        if m["redyeingPct"] is not None and m["redyeingPct"] > TARGETS["redyeing_limit"]:
            items.append({"type": "warning", "title": "Redyeing above limit",
                          "text": f"Redyeing was {m['redyeingPct']:.2f}% ({m['redyeing']:,.0f} kg) vs a "
                                  f"{TARGETS['redyeing_limit']:.1f}% limit."})

        # Loss category
        lm = dyeing["loss"]["monthly"].get(mk)
        if lm and lm["categories"]:
            top = lm["categories"][0]
            items.append({"type": "info", "title": "Largest downtime driver",
                          "text": f"{top['category']} accounted for {top['hours']:,.1f} h "
                                  f"({top['share']:.1f}% of recorded loss time)."})

        # Machines
        bulk = [x for x in dyeing["machines"] if not x["isSample"] and x["oee"] is not None]
        if bulk:
            best = max(bulk, key=lambda x: x["oee"])
            worst = min(bulk, key=lambda x: x["oee"])
            items.append({
                "type": "info",
                "title": "Machine spread",
                "text": f"{best['machine']} led with {best['oee']:.1f}% OEE while {worst['machine']} trailed at "
                        f"{worst['oee']:.1f}% — a {best['oee'] - worst['oee']:.1f} pp spread.",
            })

        # MoM production
        if prev and prev["production"]:
            chg = (m["production"] - prev["production"]) / prev["production"] * 100
            direction = "up" if chg >= 0 else "down"
            items.append({
                "type": "positive" if chg >= 0 else "warning",
                "title": f"Output {direction} month-on-month",
                "text": f"Production moved {chg:+.1f}% versus {prev['label']} "
                        f"({prev['production']:,.0f} kg → {m['production']:,.0f} kg).",
            })

        # Finishing
        fm = fin_monthly.get(mk)
        if fm and fm["oee"] is not None:
            items.append({
                "type": "positive" if fm["oee"] >= TARGETS["finishing_oee"] else "warning",
                "title": "Finishing OEE",
                "text": f"Finishing OEE was {fm['oee']:.1f}% on {fm['production']:,.0f} kg of machine throughput "
                        f"(utilization {fm['utilization']:.1f}%, efficiency {fm['efficiency']:.1f}%).",
            })

        insights[mk] = items[:9]

    return insights


# ==========================================================================
# DATA QUALITY
# ==========================================================================
def build_data_quality(dyeing, finishing):
    months = dyeing["months"]
    latest = months[-1] if months else None
    dy_days = len(dyeing["daily"])
    fin_days = len(finishing["daily"])
    rft_months = len([r for r in dyeing["rft"]["monthly"] if r["total"]])
    return {
        "dyeingRecords": sum(m["records"] for m in months),
        "dyeingDays": dy_days,
        "dyeingMachines": len([m for m in dyeing["machines"] if not m["isSample"]]),
        "sampleMachines": len([m for m in dyeing["machines"] if m["isSample"]]),
        "blankRows": dyeing["blankRows"],
        "rftMonths": rft_months,
        "finishingRecords": sum(m["records"] for m in finishing["months"]),
        "finishingDays": fin_days,
        "finishingMachines": len(finishing["machines"]),
        "latestDyeingMonth": latest["key"] if latest else None,
        "dyeingCoverage": [
            {"month": m["label"], "records": m["records"], "days": m["workingDays"]}
            for m in months
        ],
        "finishingCoverage": [
            {"month": m["label"], "records": m["records"], "days": m["workingDays"]}
            for m in finishing["months"]
        ],
        "warnings": build_quality_warnings(dyeing, finishing),
    }


def build_quality_warnings(dyeing, finishing):
    warnings = []
    # machines with no OEE
    for m in dyeing["machines"]:
        if not m["isSample"] and m["oee"] is None and m["records"]:
            warnings.append(f"{m['machine']} has {m['records']} records but no computable OEE (missing hours or capacity).")
    # machines with unusual working days
    for m in dyeing["machines"]:
        if not m["isSample"] and m["workingHours"] == 0 and m["records"] > 0:
            warnings.append(f"{m['machine']} has production records but zero working hours logged.")
    # finishing machines with no target time
    for m in finishing["machines"]:
        if (m["production"] or 0) > 0 and (m["targetTime"] or 0) == 0:
            warnings.append(f"Finishing machine '{m['machine']}' logs {m['production']:,.0f} kg but no target time — utilization cannot be computed.")
    # finishing machine naming
    names = {m["machine"] for m in finishing["machines"]}
    if "Slitting Machine" in names and "Super Slitting Machine" in names:
        warnings.append("Two slitting machines are recorded ('Super Slitting Machine' and 'Slitting Machine') — confirm whether these are the same asset.")
    return warnings[:12]


# ==========================================================================
# MAIN
# ==========================================================================
def main():
    if not DYEING_XLSX.exists():
        sys.exit(f"Dyeing workbook not found: {DYEING_XLSX}")
    if not FINISHING_XLSX.exists():
        sys.exit(f"Finishing workbook not found: {FINISHING_XLSX}")

    print("Reading dyeing workbook ...")
    dyeing = load_dyeing()
    print(f"  {len(dyeing['months'])} months, {len(dyeing['daily'])} production days, "
          f"{len(dyeing['machines'])} machines")

    print("Reading finishing workbook ...")
    finishing = load_finishing()
    print(f"  {len(finishing['months'])} months, {len(finishing['daily'])} days, "
          f"{len(finishing['machines'])} machines")

    print("Generating insights ...")
    insights = build_insights(dyeing, finishing)
    quality = build_data_quality(dyeing, finishing)

    months = dyeing["months"]
    complete = [m for m in months if not m["partial"]]
    latest_complete = complete[-1] if complete else (months[-1] if months else None)

    payload = {
        "meta": {
            "company": "Clifton Apparels Ltd.",
            "department": "Dyeing & Finishing",
            "author": "Ikramul Hasan Rakib",
            "authorRole": "Executive, Central IE Department, Clifton Group",
            "portfolio": "https://rakib-orion.vercel.app/",
            "copyright": "\u00a9 2026 Ikramul Hasan Rakib. All rights reserved.",
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "generatedAtDisplay": datetime.now().strftime("%d %b %Y, %H:%M"),
            "sourceFiles": [DYEING_XLSX.name, FINISHING_XLSX.name],
            "months": [
                {"key": m["key"], "label": m["label"], "name": m["monthName"],
                 "year": m["year"], "partial": m["partial"]}
                for m in months
            ],
            "latestMonth": months[-1]["key"] if months else None,
            "latestCompleteMonth": latest_complete["key"] if latest_complete else None,
            "dateMin": dyeing["daily"][0]["date"] if dyeing["daily"] else None,
            "dateMax": dyeing["daily"][-1]["date"] if dyeing["daily"] else None,
            "standardWorkingDays": STANDARD_WORKING_DAYS,
            "currency": "BDT",
        },
        "targets": TARGETS,
        "dyeing": dyeing,
        "finishing": finishing,
        "insights": insights,
        "dataQuality": quality,
        "notes": [
            "Dyeing KPI are computed from the 'Production Report' sheet of DYEING PERFORMANCE OVERVIEW.xlsx.",
            "OEE = Availability x Performance x Quality. Availability = operating hours / working hours; "
            "Performance = production / loading capacity; Quality = quality-pass kg / planned kg.",
            "The 'Weighted' method aggregates plant totals before taking ratios (standard OEE practice). "
            "The 'Report' method averages row-level ratios and matches the monthly PowerPoint deck.",
            "RFT is computed from the 'RFT Report' sheet; topping and redyeing percentages use planned output as the base.",
            "Finishing KPI are computed from the 'Database' sheet of FINISHING PERFORMANCE OVERVIEW.xlsx.",
            "Sample dyeing machines (Mc-1, Mc-2, Mc-3, Mc-25, Mc-26) are excluded from machine rankings.",
            "Dashboard designed and developed by Ikramul Hasan Rakib, Executive, Central IE Department, "
            "Clifton Group. (c) 2026 Ikramul Hasan Rakib. All rights reserved. "
            "Proprietary license - see the LICENSE and NOTICE files in the project root.",
        ],
    }

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    banner = (
        "/* Dyeing & Finishing Performance Dashboard - data bundle\n"
        "   Author : Ikramul Hasan Rakib\n"
        "            Executive, Central IE Department, Clifton Group\n"
        "   (c) 2026 Ikramul Hasan Rakib. All rights reserved.\n"
        "   Proprietary license - see LICENSE and NOTICE in the project root.\n"
        "   Generated automatically by scripts/build.py - do not edit by hand. */\n"
    )
    js = banner + "window.DASHBOARD_DATA = " + json.dumps(payload, ensure_ascii=True, separators=(",", ":")) + ";\n"
    (OUT_DIR / "data.js").write_text(js, encoding="utf-8")
    (OUT_DIR / "data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8"
    )

    size_kb = (OUT_DIR / "data.js").stat().st_size / 1024
    print(f"\nWrote dashboard/data/data.js  ({size_kb:,.0f} KB)")
    print(f"Wrote dashboard/data/data.json")
    print(f"Latest month: {payload['meta']['latestMonth']}  |  latest complete: {payload['meta']['latestCompleteMonth']}")
    print("Done.")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        sys.exit(1)
